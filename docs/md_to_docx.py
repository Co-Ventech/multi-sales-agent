import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

MD_FILE = r"G:\company-agents\Sales-outreach-agent-v2\docs\agent-overview.md"
OUT_FILE = r"G:\company-agents\Sales-outreach-agent-v2\docs\agent-overview.docx"

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

# Styles
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)

def set_heading(para, level):
    sizes = {1: 20, 2: 15, 3: 13, 4: 11}
    para.style = doc.styles[f'Heading {level}']
    run = para.runs[0] if para.runs else para.add_run(para.text)
    run.font.size = Pt(sizes.get(level, 11))
    if level == 1:
        run.font.color.rgb = RGBColor(0x0f, 0x34, 0x60)
    elif level == 2:
        run.font.color.rgb = RGBColor(0x19, 0x87, 0x54)
    else:
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_inline_text(para, text):
    """Parse **bold**, `code`, and plain text inline."""
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xc7, 0x25, 0x4e)
        else:
            if part:
                para.add_run(part)

with open(MD_FILE, encoding='utf-8') as f:
    lines = f.readlines()

i = 0
while i < len(lines):
    raw = lines[i].rstrip('\n')

    # Skip empty lines
    if raw.strip() == '':
        i += 1
        continue

    # Headings
    m = re.match(r'^(#{1,4})\s+(.*)', raw)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        para = doc.add_paragraph()
        para.clear()
        run = para.add_run(text)
        set_heading(para, level)
        i += 1
        continue

    # Horizontal rule
    if re.match(r'^---+$', raw.strip()):
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run('─' * 80)
        run.font.color.rgb = RGBColor(0xcc, 0xcc, 0xcc)
        run.font.size = Pt(7)
        i += 1
        continue

    # Table
    if raw.strip().startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].rstrip('\n'))
            i += 1
        # Filter out separator rows (---|---)
        data_rows = [r for r in table_lines if not re.match(r'^\s*\|[-| :]+\|\s*$', r)]
        if data_rows:
            cols = [c.strip() for c in data_rows[0].split('|')[1:-1]]
            num_cols = len(cols)
            tbl = doc.add_table(rows=1, cols=num_cols)
            tbl.style = 'Table Grid'
            # Header row
            hdr = tbl.rows[0].cells
            for ci, col in enumerate(cols):
                hdr[ci].text = col
                for run in hdr[ci].paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9.5)
                hdr[ci].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xff, 0xff, 0xff) if True else None
                # Header bg color via XML
                from docx.oxml.ns import qn
                from docx.oxml import OxmlElement
                tc = hdr[ci]._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), '0f3460')
                tcPr.append(shd)
                for run in hdr[ci].paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
            # Data rows
            for row_text in data_rows[1:]:
                cells = [c.strip() for c in row_text.split('|')[1:-1]]
                row = tbl.add_row()
                for ci, cell_text in enumerate(cells):
                    if ci < len(row.cells):
                        para = row.cells[ci].paragraphs[0]
                        add_inline_text(para, cell_text)
                        para.runs and setattr(para.runs[0], 'font_size', Pt(9.5))
        doc.add_paragraph()
        continue

    # Code block
    if raw.strip().startswith('```'):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code_lines.append(lines[i].rstrip('\n'))
            i += 1
        i += 1
        para = doc.add_paragraph()
        para.style = 'Normal'
        para.paragraph_format.left_indent = Inches(0.3)
        run = para.add_run('\n'.join(code_lines))
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pPr = para._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'f4f4f4')
        pPr.append(shd)
        continue

    # Blockquote
    if raw.strip().startswith('>'):
        text = raw.strip().lstrip('> ').strip()
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.4)
        add_inline_text(para, text)
        for run in para.runs:
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            run.italic = True
        i += 1
        continue

    # Bullet list
    if re.match(r'^[-*]\s+', raw.strip()):
        text = re.sub(r'^[-*]\s+', '', raw.strip())
        para = doc.add_paragraph(style='List Bullet')
        add_inline_text(para, text)
        i += 1
        continue

    # Numbered list
    m = re.match(r'^\d+\.\s+(.*)', raw.strip())
    if m:
        para = doc.add_paragraph(style='List Number')
        add_inline_text(para, m.group(1))
        i += 1
        continue

    # Regular paragraph
    para = doc.add_paragraph()
    add_inline_text(para, raw.strip())
    i += 1

doc.save(OUT_FILE)
print(f"Saved: {OUT_FILE}")
